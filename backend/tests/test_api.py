"""End-to-end API tests covering the B3 acceptance scenarios.

Run from backend/:
    py -m pytest -q
"""
import csv
import io
from datetime import date
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.main import app

BACKEND_DIR = Path(__file__).resolve().parents[1]
SAMPLES_DIR = BACKEND_DIR.parent / "samples"
DIST_INDEX = BACKEND_DIR.parent / "frontend" / "dist" / "index.html"


def _teacher_payload(teacher_no: str, **overrides) -> dict:
    payload = {
        "teacher_no": teacher_no,
        "name": f"测试教师{teacher_no[-3:]}",
        "gender": "女",
        "birth_year": 1985,
        "college": "测试学院",
        "department": "测试系",
        "discipline": "测试学科",
        "degree": "博士",
        "title": "副教授",
        "position": "专任教师",
        "hire_date": "2012-09-01",
        "research_direction": "测试方向 数据治理",
        "bio": "自动化测试教师档案。",
        "skills": "数据分析、系统设计",
        "research_keywords": "数据治理,测试",
        "status": "active",
        "advisor_qualified": False,
        "enrollment_status": "closed",
        "capacity": 0,
        "current_students": 0,
        "phone": "13800138000",
        "email": "teacher@example.com",
        "data_source": "自动化测试",
        "verification_status": "verified",
    }
    payload.update(overrides)
    return payload


def _create_teacher(client, headers, teacher_no: str, **overrides) -> dict:
    response = client.post(
        "/api/teachers", json=_teacher_payload(teacher_no, **overrides), headers=headers
    )
    assert response.status_code == 200, response.text
    return response.json()


def _csv_bytes(headers: list[str], rows: list[dict]) -> bytes:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=headers, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8-sig")


def _find(items: list[dict], **conditions) -> dict:
    for item in items:
        if all(item.get(key) == value for key, value in conditions.items()):
            return item
    raise AssertionError(f"未找到满足条件的记录：{conditions}")


# --------------------------------------------------------------------------
# 1. 健康检查、前端页面与单服务部署
# --------------------------------------------------------------------------
def test_health_and_frontend_routes(client):
    health = client.get("/api/health")
    assert health.status_code == 200
    data = health.json()
    assert data["status"] == "ok"
    assert data["database"] == "ok"
    assert data["app"].startswith("B3")

    if not DIST_INDEX.exists():
        pytest.skip("frontend/dist 尚未构建，跳过前端静态页面测试")
    for path in ("/", "/login", "/teachers", "/advisor"):
        response = client.get(path)
        assert response.status_code == 200, path
        assert response.headers["content-type"].startswith("text/html")
        assert '<div id="app"></div>' in response.text

    unknown = client.get("/api/does-not-exist")
    assert unknown.status_code == 404
    assert unknown.headers["content-type"].startswith("application/json")
    assert "接口不存在" in unknown.json()["detail"]


# --------------------------------------------------------------------------
# 2. 登录、认证与中文错误提示
# --------------------------------------------------------------------------
def test_auth_and_error_messages(client):
    anonymous = client.get("/api/teachers")
    assert anonymous.status_code == 401
    assert anonymous.json()["detail"] == "请先登录"

    bad = client.post("/api/auth/login", json={"username": "admin", "password": "bad"})
    assert bad.status_code == 401
    assert "用户名或密码错误" in bad.json()["detail"]

    login = client.post(
        "/api/auth/login", json={"username": "manager", "password": "manager123"}
    )
    assert login.status_code == 200
    body = login.json()
    assert body["token_type"] == "bearer"
    assert body["user"]["role"] == "manager"
    assert body["user"]["role_label"] == "业务人员"

    headers = {"Authorization": f"Bearer {body['access_token']}"}
    me = client.get("/api/auth/me", headers=headers)
    assert me.status_code == 200
    assert me.json()["username"] == "manager"


# --------------------------------------------------------------------------
# 3. 教师档案 CRUD、搜索与分页
# --------------------------------------------------------------------------
def test_teacher_crud_search_and_pagination(client, manager_headers):
    created = _create_teacher(client, manager_headers, "TEST-CRUD-001")
    teacher_id = created["id"]
    assert created["verification_status"] == "verified"
    assert created["available_capacity"] == 0
    assert created["age"] == date.today().year - 1985

    options = client.get("/api/teachers/options", headers=manager_headers).json()
    assert options["disciplines_by_college"]["测试学院"] == ["测试学科"]

    updated = client.put(
        f"/api/teachers/{teacher_id}",
        json={
            "title": "教授",
            "capacity": 5,
            "current_students": 2,
            "research_direction": "数据治理 知识图谱",
        },
        headers=manager_headers,
    )
    assert updated.status_code == 200, updated.text
    updated_body = updated.json()
    assert updated_body["title"] == "教授"
    assert updated_body["available_capacity"] == 3

    search = client.get(
        "/api/teachers",
        params={"keyword": "TEST-CRUD-001", "page": 1, "page_size": 10},
        headers=manager_headers,
    )
    assert search.status_code == 200
    assert search.json()["total"] == 1
    assert search.json()["items"][0]["id"] == teacher_id

    detail = client.get(f"/api/teachers/{teacher_id}", headers=manager_headers)
    assert detail.status_code == 200
    assert detail.json()["counts"] == {
        "outputs": 0, "teaching": 0, "services": 0, "projects": 0
    }

    deleted = client.delete(f"/api/teachers/{teacher_id}", headers=manager_headers)
    assert deleted.status_code == 200
    assert client.get(f"/api/teachers/{teacher_id}", headers=manager_headers).status_code == 404


# --------------------------------------------------------------------------
# 4. 关联数据保护：有成果的教师不可直接删除
# --------------------------------------------------------------------------
def test_teacher_delete_is_blocked_by_related_records(client, manager_headers):
    teacher = _create_teacher(client, manager_headers, "TEST-DEP-001")
    output = client.post(
        "/api/research-outputs",
        json={
            "teacher_id": teacher["id"],
            "output_type": "paper",
            "title": "关联保护测试论文",
            "publish_date": "2024-01-01",
        },
        headers=manager_headers,
    )
    assert output.status_code == 200, output.text

    blocked = client.delete(f"/api/teachers/{teacher['id']}", headers=manager_headers)
    assert blocked.status_code == 409
    assert "科研成果" in blocked.json()["detail"]
    assert "不能删除" in blocked.json()["detail"]

    assert client.delete(
        f"/api/research-outputs/{output.json()['id']}", headers=manager_headers
    ).status_code == 200
    assert client.delete(
        f"/api/teachers/{teacher['id']}", headers=manager_headers
    ).status_code == 200


# --------------------------------------------------------------------------
# 5. 参数校验与重复约束
# --------------------------------------------------------------------------
def test_teacher_validation_and_duplicate(client, manager_headers):
    invalid = client.post(
        "/api/teachers",
        json=_teacher_payload("TEST-BAD-001", title="首席教授", email="not-an-email"),
        headers=manager_headers,
    )
    assert invalid.status_code == 422
    detail = invalid.json()["detail"]
    assert detail.startswith("参数校验失败")
    assert "职称" in detail and "邮箱" in detail
    assert invalid.json()["errors"]

    capacity = client.post(
        "/api/teachers",
        json=_teacher_payload("TEST-BAD-002", capacity=1, current_students=2),
        headers=manager_headers,
    )
    assert capacity.status_code == 400
    assert "指导容量不能小于当前在读人数" in capacity.json()["detail"]

    duplicate = client.post(
        "/api/teachers",
        json=_teacher_payload("B3T0001", name="重复编号测试"),
        headers=manager_headers,
    )
    assert duplicate.status_code == 400
    assert "教师编号已存在" in duplicate.json()["detail"]


# --------------------------------------------------------------------------
# 6. Excel/CSV 导入：成功、失败、重复与错误导出
# --------------------------------------------------------------------------
def test_teacher_and_output_imports(client, manager_headers):
    headers = ["teacher_no", "name", "college", "discipline", "title", "email"]
    rows = [
        {
            "teacher_no": "TEST-IMP-001",
            "name": "导入教师甲",
            "college": "测试学院",
            "discipline": "测试学科",
            "title": "讲师",
            "email": "imp1@example.com",
        },
        {
            "teacher_no": "B3T0001",
            "name": "重复编号教师",
            "college": "测试学院",
            "discipline": "测试学科",
            "title": "讲师",
            "email": "dup@example.com",
        },
        {
            "teacher_no": "TEST-IMP-003",
            "name": "错误教师",
            "college": "测试学院",
            "discipline": "测试学科",
            "title": "首席教授",
            "email": "bad-email",
        },
    ]
    response = client.post(
        "/api/imports/teachers",
        files={"file": ("teachers_test.csv", _csv_bytes(headers, rows), "text/csv")},
        headers=manager_headers,
    )
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["success_rows"] == 1
    assert result["failed_rows"] == 1
    assert result["duplicate_rows"] == 1
    assert result["total_rows"] == 3
    assert result["status"] == "completed"

    batches = client.get(
        "/api/imports/batches", params={"import_type": "teachers"}, headers=manager_headers
    )
    assert batches.status_code == 200
    batch = _find(batches.json()["items"], id=result["id"])
    assert batch["error_count"] >= 2

    errors = client.get(
        f"/api/imports/batches/{result['id']}/errors", headers=manager_headers
    )
    assert errors.status_code == 200
    error_text = " ".join(item["error_message"] for item in errors.json()["items"])
    assert "重复" in error_text
    assert "职称不合法" in error_text or "邮箱格式" in error_text

    exported = client.get(
        f"/api/imports/batches/{result['id']}/errors/export", headers=manager_headers
    )
    assert exported.status_code == 200
    assert exported.content.startswith(b"\xef\xbb\xbf")
    assert "行号,字段,错误原因,原始值" in exported.content.decode("utf-8-sig")

    imported = client.get(
        "/api/teachers",
        params={"keyword": "TEST-IMP-001"},
        headers=manager_headers,
    ).json()["items"][0]
    assert imported["verification_status"] == "pending"
    assert imported["data_source"].startswith("导入:")

    # Excel 路径
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    sheet.append([
        "TEST-IMP-XLSX-001", "Excel导入教师", "测试学院", "测试学科", "讲师", "xlsx@example.com"
    ])
    stream = io.BytesIO()
    workbook.save(stream)
    excel = client.post(
        "/api/imports/teachers",
        files={
            "file": (
                "teachers_test.xlsx",
                stream.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=manager_headers,
    )
    assert excel.status_code == 200, excel.text
    assert excel.json()["success_rows"] == 1

    # 科研成果导入：一行成功、一行教师不存在、一行同文件重复
    output_headers = ["teacher_no", "output_type", "title", "publish_date"]
    output_rows = [
        {
            "teacher_no": "TEST-IMP-001",
            "output_type": "论文",
            "title": "导入测试成果",
            "publish_date": "2023-05-01",
        },
        {
            "teacher_no": "NOT-EXIST-999",
            "output_type": "论文",
            "title": "教师不存在成果",
            "publish_date": "2023-05-01",
        },
        {
            "teacher_no": "TEST-IMP-001",
            "output_type": "论文",
            "title": "导入测试成果",
            "publish_date": "2023-09-01",
        },
    ]
    output_import = client.post(
        "/api/imports/research-outputs",
        files={
            "file": (
                "outputs_test.csv",
                _csv_bytes(output_headers, output_rows),
                "text/csv",
            )
        },
        headers=manager_headers,
    )
    assert output_import.status_code == 200, output_import.text
    out = output_import.json()
    assert out["success_rows"] == 1
    assert out["failed_rows"] == 1
    assert out["duplicate_rows"] == 1

    outputs = client.get(
        "/api/research-outputs",
        params={"keyword": "导入测试成果"},
        headers=manager_headers,
    )
    assert outputs.status_code == 200
    assert outputs.json()["total"] == 1


def test_demo_sample_files_are_importable(client, manager_headers):
    teacher_file = SAMPLES_DIR / "teachers_import_demo.csv"
    output_file = SAMPLES_DIR / "outputs_import_demo.csv"
    candidate_file = SAMPLES_DIR / "candidates_import_demo.csv"
    assert teacher_file.exists() and output_file.exists() and candidate_file.exists()

    teachers = client.post(
        "/api/imports/teachers",
        files={"file": (teacher_file.name, teacher_file.read_bytes(), "text/csv")},
        headers=manager_headers,
    )
    assert teachers.status_code == 200, teachers.text
    teacher_result = teachers.json()
    assert (teacher_result["success_rows"], teacher_result["failed_rows"],
            teacher_result["duplicate_rows"]) == (1, 1, 1)

    outputs = client.post(
        "/api/imports/research-outputs",
        files={"file": (output_file.name, output_file.read_bytes(), "text/csv")},
        headers=manager_headers,
    )
    assert outputs.status_code == 200, outputs.text
    output_result = outputs.json()
    assert (output_result["success_rows"], output_result["failed_rows"],
            output_result["duplicate_rows"]) == (1, 1, 1)

    positions = client.get(
        "/api/recruitment/positions", headers=manager_headers
    ).json()["items"]
    assert positions
    candidates = client.post(
        "/api/imports/candidates",
        params={"position_id": positions[0]["id"]},
        files={"file": (candidate_file.name, candidate_file.read_bytes(), "text/csv")},
        headers=manager_headers,
    )
    assert candidates.status_code == 200, candidates.text
    candidate_result = candidates.json()
    assert candidate_result["success_rows"] == 1
    assert candidate_result["duplicate_rows"] == 1


# --------------------------------------------------------------------------
# 7. 画像统计与数据库记录一致
# --------------------------------------------------------------------------
def test_profile_statistics_match_database(client, admin_headers, viewer_headers):
    first_output = client.get(
        "/api/research-outputs", params={"page_size": 1}, headers=admin_headers
    ).json()["items"][0]
    teacher_id = first_output["teacher_id"]

    profile = client.get(f"/api/teachers/{teacher_id}/profile", headers=admin_headers)
    assert profile.status_code == 200, profile.text
    data = profile.json()
    assert data["output_count"] > 0
    assert sum(data["output_by_type"].values()) == data["output_count"]
    assert data["project_count"] >= data["principal_project_count"]
    assert data["teaching_record_count"] >= 0
    assert data["service_count"] == (
        data["award_count"] + data["academic_post_count"] + data["social_service_count"]
    )
    missing = [item for item in data["completeness_detail"] if not item["present"]]
    assert len(missing) == len(data["missing_fields"])
    assert data["completeness"] == round(
        (11 - len(data["missing_fields"])) / 11 * 100, 1
    )
    assert len(data["outputs"]) == min(20, data["output_count"])
    assert data["metric_definitions"] and data["local_summary"]
    assert data["llm_status"]["mode"] == "local"
    assert data["sensitive_hidden"] is False

    outputs = client.get(
        "/api/research-outputs",
        params={"teacher_id": teacher_id, "page_size": 100},
        headers=admin_headers,
    ).json()
    assert outputs["total"] == data["output_count"]

    teaching = client.get(
        "/api/teaching-records",
        params={"teacher_id": teacher_id, "page_size": 100},
        headers=admin_headers,
    ).json()
    assert teaching["total"] == data["teaching_record_count"]
    assert round(sum(item["hours"] or 0 for item in teaching["items"]), 1) == data["teaching_hours_total"]

    viewer_profile = client.get(
        f"/api/teachers/{teacher_id}/profile", headers=viewer_headers
    )
    assert viewer_profile.status_code == 200
    assert viewer_profile.json()["sensitive_hidden"] is True
    assert "phone" not in viewer_profile.json()["teacher"]
    assert "email" not in viewer_profile.json()["teacher"]


# --------------------------------------------------------------------------
# 8. 人才驾驶舱、盘点筛选、人才规划缺口与 Excel 导出
# --------------------------------------------------------------------------
def test_dashboard_talent_filters_plan_and_export(client, admin_headers):
    options = client.get("/api/analytics/filters", headers=admin_headers).json()
    assert options["colleges"] and options["disciplines"]
    mapping = options["disciplines_by_college"]
    assert mapping and set(mapping) == set(options["colleges"])
    assert all(set(disciplines) <= set(options["disciplines"]) for disciplines in mapping.values())

    overall = client.get("/api/analytics/dashboard", headers=admin_headers).json()
    assert overall["total_teachers"] > 0
    assert overall["active_teachers"] <= overall["total_teachers"]
    assert sum(item["value"] for item in overall["college_distribution"]) == overall["total_teachers"]
    assert sum(item["value"] for item in overall["title_distribution"]) == overall["total_teachers"]
    assert overall["stats_scope"] and overall["age_note"]
    assert "data_quality" in overall

    college = options["colleges"][0]
    expected = _find(overall["college_distribution"], name=college)["value"]
    filtered = client.get(
        "/api/analytics/dashboard", params={"college": college}, headers=admin_headers
    ).json()
    assert filtered["total_teachers"] == expected
    assert [item["name"] for item in filtered["college_distribution"]] == [college]
    assert filtered["output_total"] <= overall["output_total"]
    assert all(item["year"] for item in filtered["output_by_year"])

    review = client.get(
        "/api/analytics/talent-review", params={"college": college}, headers=admin_headers
    )
    assert review.status_code == 200
    review_data = review.json()
    assert review_data["overview"]["total_teachers"] == expected
    assert len(review_data["title_structure_by_college"]) == 1
    assert review_data["title_structure_by_college"][0]["total"] == expected
    assert review_data["plan_note"]

    # 管理员配置人才规划后，缺口来自配置而非模型猜测
    discipline = filtered["discipline_distribution"][0]["name"]
    target = expected + 7
    try:
        plan = client.put(
            "/api/settings/talent-plan",
            json={"items": [{
                "college": college,
                "discipline": discipline,
                "target_count": target,
                "note": "自动化测试规划",
            }]},
            headers=admin_headers,
        )
        assert plan.status_code == 200
        review_with_plan = client.get(
            "/api/analytics/talent-review",
            params={"college": college, "discipline": discipline},
            headers=admin_headers,
        ).json()
        gap = _find(review_with_plan["plan_gaps"], college=college, discipline=discipline)
        assert gap["current_count"] <= expected
        assert gap["target_count"] == target
        assert gap["gap"] == target - gap["current_count"]
    finally:
        client.put(
            "/api/settings/talent-plan", json={"items": []}, headers=admin_headers
        )

    export = client.get(
        "/api/analytics/talent-review/export",
        params={"college": college},
        headers=admin_headers,
    )
    assert export.status_code == 200
    assert "spreadsheetml" in export.headers["content-type"]
    book = load_workbook(io.BytesIO(export.content))
    assert "人才盘点概览" in book.sheetnames
    assert "学院职称结构" in book.sheetnames
    labels = [cell.value for cell in book["人才盘点概览"]["A"]]
    assert "教师总人数" in labels
    assert "统计说明" in labels


# --------------------------------------------------------------------------
# 9. 数据质量检查
# --------------------------------------------------------------------------
def test_data_quality_endpoint(client, admin_headers):
    quality = client.get("/api/analytics/data-quality", headers=admin_headers)
    assert quality.status_code == 200
    data = quality.json()
    assert len(data["missing_field_counts"]) == 7
    assert data["duplicate_teacher_count"] >= len(data["duplicate_teachers"])
    assert data["duplicate_output_count"] >= 0
    assert data["pending_verification_count"] >= 1  # 导入的数据处于待核验状态
    assert "不代表成果真实性" in data["note"]


# --------------------------------------------------------------------------
# 10. 导师匹配：硬性条件优先、可解释评分、人工确认
# --------------------------------------------------------------------------
def test_advisor_matching_hard_filters_and_explanations(client, manager_headers):
    students = client.get(
        "/api/students", params={"page_size": 100}, headers=manager_headers
    ).json()["items"]
    assert students
    student = next(
        (s for s in students if s.get("desired_direction") or s.get("research_interest")),
        students[0],
    )
    direction = student.get("desired_direction") or student.get("research_interest") or "测试方向"
    keywords = student.get("keywords") or direction
    needs = student.get("training_needs") or "科研训练"

    ineligible = [
        _create_teacher(
            client, manager_headers, "TEST-ADV-Q0",
            college="测试学院", discipline=student["discipline"], title="教授",
            research_direction=direction, research_keywords=keywords, skills=needs,
            bio=needs, advisor_qualified=False, enrollment_status="open",
            capacity=10, current_students=0,
        ),
        _create_teacher(
            client, manager_headers, "TEST-ADV-EN",
            college="测试学院", discipline=student["discipline"], title="教授",
            research_direction=direction, research_keywords=keywords, skills=needs,
            bio=needs, advisor_qualified=True, enrollment_status="closed",
            capacity=10, current_students=0,
        ),
        _create_teacher(
            client, manager_headers, "TEST-ADV-CAP",
            college="测试学院", discipline=student["discipline"], title="教授",
            research_direction=direction, research_keywords=keywords, skills=needs,
            bio=needs, advisor_qualified=True, enrollment_status="open",
            capacity=3, current_students=3,
        ),
    ]

    bad_weights = client.post(
        "/api/advisor/tasks",
        json={"student_id": student["id"], "weights": {"unexpected": 1}},
        headers=manager_headers,
    )
    assert bad_weights.status_code == 400
    assert "未知权重维度" in bad_weights.json()["detail"]

    task = client.post(
        "/api/advisor/tasks",
        json={
            "student_id": student["id"],
            "name": "自动化测试导师匹配",
            "weights": {"direction": 1, "topic": 1, "needs": 1, "capacity": 1},
        },
        headers=manager_headers,
    )
    assert task.status_code == 200, task.text
    task_id = task.json()["id"]
    run = client.post(f"/api/advisor/tasks/{task_id}/run", headers=manager_headers)
    assert run.status_code == 200, run.text
    summary = run.json()["summary"]
    assert summary["qualified_count"] + summary["excluded_count"] == summary["candidate_teachers"]
    assert summary["qualified_count"] > 0
    assert summary["excluded_count"] > 0

    result = client.get(f"/api/advisor/tasks/{task_id}", headers=manager_headers)
    assert result.status_code == 200
    data = result.json()
    recommendations = data["recommendations"]
    excluded = data["excluded"]
    assert len(recommendations) == summary["qualified_count"]
    assert len(excluded) == summary["excluded_count"]
    assert [row["rank"] for row in recommendations] == list(range(1, len(recommendations) + 1))
    scores = [row["total_score"] for row in recommendations]
    assert scores == sorted(scores, reverse=True)
    assert all(row["qualified"] and not row["disqualifiers"] for row in recommendations)
    assert all(not row["qualified"] and row["disqualifiers"] for row in excluded)
    assert all(row["rank"] is None for row in excluded)
    assert "研究方向相似度" in recommendations[0]["reasons"]
    assert "评分权重" in recommendations[0]["reasons"]
    assert abs(sum(data["weights"].values()) - 1.0) < 0.001
    assert data["summary"]["note"].startswith("仅满足导师资格")

    excluded_by_id = {row["teacher_id"]: row for row in excluded}
    qualification = excluded_by_id[ineligible[0]["id"]]
    enrollment = excluded_by_id[ineligible[1]["id"]]
    capacity = excluded_by_id[ineligible[2]["id"]]
    assert "未取得导师指导资格" in qualification["disqualifiers"]
    assert "招生状态未开放" in enrollment["disqualifiers"]
    assert "指导容量已满" in capacity["disqualifiers"]
    assert qualification["total_score"] > 0  # 即使规则分不低，也不能进入推荐
    recommended_ids = {row["teacher_id"] for row in recommendations}
    assert not recommended_ids.intersection({row["id"] for row in ineligible})

    confirm = client.put(
        f"/api/advisor/results/{recommendations[0]['id']}/confirm",
        json={"human_confirmed": True, "note": "自动化测试人工确认"},
        headers=manager_headers,
    )
    assert confirm.status_code == 200
    assert confirm.json()["human_confirmed"] is True
    refreshed = client.get(f"/api/advisor/tasks/{task_id}", headers=manager_headers).json()
    assert refreshed["recommendations"][0]["note"] == "自动化测试人工确认"


# --------------------------------------------------------------------------
# 10b. 师生匹配度：0-100 量化、分项贡献、等级与硬条件独立校验
# --------------------------------------------------------------------------
def test_match_degree_quantification(client, manager_headers, viewer_headers):
    students = client.get(
        "/api/students", params={"page_size": 100}, headers=manager_headers
    ).json()["items"]
    student = next(
        (s for s in students if s.get("desired_direction") or s.get("research_interest")),
        students[0],
    )
    direction = student.get("desired_direction") or student.get("research_interest") or "测试方向"
    keywords = student.get("keywords") or direction
    needs = student.get("training_needs") or "科研训练"

    teacher = _create_teacher(
        client, manager_headers, "TEST-MATCH-001",
        college="测试学院", discipline=student["discipline"], title="教授",
        research_direction=direction, research_keywords=keywords,
        skills=needs, bio=needs, advisor_qualified=True,
        enrollment_status="open", capacity=5, current_students=0,
    )
    weights = {"direction": 1, "topic": 1, "needs": 1, "capacity": 1}
    single = client.post(
        "/api/advisor/score",
        json={"student_id": student["id"], "teacher_id": teacher["id"], "weights": weights},
        headers=manager_headers,
    )
    assert single.status_code == 200, single.text
    data = single.json()
    assert data["scope"] == "single"
    assert data["student"]["id"] == student["id"]
    assert [rule["grade"] for rule in data["grade_rules"]] == ["A", "B", "C", "D"]
    assert len(data["items"]) == 1
    item = data["items"][0]
    assert item["teacher_id"] == teacher["id"]
    assert item["qualified"] is True
    assert item["rank"] >= 1
    assert 0 <= item["total_score"] <= 100
    assert item["grade"] in {"A", "B", "C", "D"}
    assert set(item["dimensions"]) == {"direction", "topic", "needs", "capacity"}
    assert abs(sum(dim["contribution"] for dim in item["dimensions"].values()) - item["total_score"]) <= 0.5
    assert item["formula"].startswith("总分 ")
    assert "研究方向相似度" in item["reasons"]
    assert "评分权重" in item["reasons"]
    assert data["summary"]["teacher_count"] == 1
    assert data["summary"]["qualified_count"] == 1
    assert data["summary"]["highest_qualified_score"] == item["total_score"]

    all_teachers = client.post(
        "/api/advisor/score",
        json={"student_id": student["id"], "weights": weights},
        headers=manager_headers,
    )
    assert all_teachers.status_code == 200, all_teachers.text
    all_data = all_teachers.json()
    assert all_data["scope"] == "all"
    summary = all_data["summary"]
    assert summary["teacher_count"] == summary["qualified_count"] + summary["excluded_count"]
    assert summary["teacher_count"] >= 100
    assert any(row["teacher_id"] == teacher["id"] for row in all_data["items"])
    qualified_flags = [row["qualified"] for row in all_data["items"]]
    assert qualified_flags == sorted(qualified_flags, reverse=True)
    qualified_scores = [row["total_score"] for row in all_data["items"] if row["qualified"]]
    assert qualified_scores == sorted(qualified_scores, reverse=True)

    viewer_view = client.post(
        "/api/advisor/score",
        json={"student_id": student["id"], "teacher_id": teacher["id"]},
        headers=viewer_headers,
    )
    assert viewer_view.status_code == 200
    viewer_teacher = viewer_view.json()["items"][0]["teacher"]
    assert viewer_teacher["phone_masked"] is True
    assert viewer_teacher["phone"] == "138****8000"

    missing_student = client.post(
        "/api/advisor/score", json={"student_id": 999999}, headers=manager_headers
    )
    assert missing_student.status_code == 404
    assert "学生档案不存在" in missing_student.json()["detail"]
    missing_teacher = client.post(
        "/api/advisor/score",
        json={"student_id": student["id"], "teacher_id": 999999},
        headers=manager_headers,
    )
    assert missing_teacher.status_code == 404
    assert "教师档案不存在" in missing_teacher.json()["detail"]
    bad_weights = client.post(
        "/api/advisor/score",
        json={"student_id": student["id"], "weights": {"unknown": 1}},
        headers=manager_headers,
    )
    assert bad_weights.status_code == 400
    assert "未知权重维度" in bad_weights.json()["detail"]


# --------------------------------------------------------------------------
# 11. 职称申报：材料归集、缺项/重复/冲突/待核验、人工确认与导出
# --------------------------------------------------------------------------
def test_promotion_gather_checks_and_export(client, manager_headers):
    teacher = _create_teacher(client, manager_headers, "TEST-PROM-001")
    output_payloads = [
        {"title": "智能测试成果A", "publish_date": "2022-06-01"},
        {"title": "智能测试成果A", "publish_date": "2022-07-01"},
        {"title": "智能测试成果B", "publish_date": None},
        {"title": "智能测试成果C", "publish_date": "2024-06-01"},
    ]
    for payload in output_payloads:
        response = client.post(
            "/api/research-outputs",
            json={
                "teacher_id": teacher["id"],
                "output_type": "paper",
                **payload,
            },
            headers=manager_headers,
        )
        assert response.status_code == 200, response.text

    batch = client.post(
        "/api/promotion/batches",
        json={
            "name": "自动化测试职称批次",
            "year": 2026,
            "title_type": "副教授",
            "college": "测试学院",
            "deadline": "2023-12-31",
            "status": "open",
        },
        headers=manager_headers,
    )
    assert batch.status_code == 200, batch.text
    batch_id = batch.json()["id"]

    requirement_1 = client.post(
        f"/api/promotion/batches/{batch_id}/requirements",
        json={
            "category": "paper",
            "name": "测试论文要求",
            "requirement_desc": "近十年智能测试成果",
            "min_count": 1,
            "mandatory": True,
            "keywords": "智能测试成果",
            "window_years": 10,
            "principal_only": False,
        },
        headers=manager_headers,
    )
    assert requirement_1.status_code == 200, requirement_1.text
    requirement_2 = client.post(
        f"/api/promotion/batches/{batch_id}/requirements",
        json={
            "category": "paper",
            "name": "测试缺失要求",
            "min_count": 1,
            "mandatory": True,
            "keywords": "ZZZ不存在关键词",
            "window_years": 10,
        },
        headers=manager_headers,
    )
    assert requirement_2.status_code == 200, requirement_2.text

    detail = client.get(f"/api/promotion/batches/{batch_id}", headers=manager_headers)
    assert detail.status_code == 200
    assert len(detail.json()["requirements"]) == 2
    assert "不代表真实高校政策" in detail.json()["disclaimer"]

    gather = client.post(
        f"/api/promotion/batches/{batch_id}/gather",
        json={"teacher_id": teacher["id"]},
        headers=manager_headers,
    )
    assert gather.status_code == 200, gather.text
    gather_result = gather.json()["results"][0]
    assert gather_result["created"] >= 5
    assert gather_result["missing_mandatory"] >= 1

    materials = client.get(
        "/api/promotion/materials",
        params={"batch_id": batch_id, "teacher_id": teacher["id"]},
        headers=manager_headers,
    )
    assert materials.status_code == 200
    items = materials.json()["items"]
    statuses = {item["status"] for item in items}
    assert {"matched", "duplicate", "pending_verification", "conflict", "missing"} <= statuses
    assert all(item["status_label"] for item in items)
    assert all(item["requirement_name"] for item in items)

    summary = client.get(
        f"/api/promotion/batches/{batch_id}/summary",
        params={"teacher_id": teacher["id"]},
        headers=manager_headers,
    )
    assert summary.status_code == 200
    summary_data = summary.json()
    assert summary_data["completion"] == 50.0
    assert len(summary_data["missing_mandatory"]) == 1
    assert summary_data["missing_mandatory"][0]["name"] == "测试缺失要求"
    assert "不验证成果真伪" in summary_data["disclaimer"]

    missing_material = _find(items, status="missing")
    invalid = client.put(
        f"/api/promotion/materials/{missing_material['id']}",
        json={"status": "unknown"},
        headers=manager_headers,
    )
    assert invalid.status_code == 400
    assert "材料状态不合法" in invalid.json()["detail"]

    confirmed = client.put(
        f"/api/promotion/materials/{missing_material['id']}",
        json={"status": "confirmed", "note": "人工核实后补充说明"},
        headers=manager_headers,
    )
    assert confirmed.status_code == 200
    assert confirmed.json()["status"] == "confirmed"
    assert confirmed.json()["updated_by"] == "manager"

    summary_after = client.get(
        f"/api/promotion/batches/{batch_id}/summary",
        params={"teacher_id": teacher["id"]},
        headers=manager_headers,
    ).json()
    assert summary_after["completion"] == 100.0
    assert summary_after["missing_mandatory"] == []

    export = client.get(
        f"/api/promotion/batches/{batch_id}/export",
        params={"teacher_id": teacher["id"]},
        headers=manager_headers,
    )
    assert export.status_code == 200
    book = load_workbook(io.BytesIO(export.content))
    assert {"申报汇总", "要求达成情况", "材料明细"} <= set(book.sheetnames)
    summary_labels = [cell.value for cell in book["申报汇总"]["A"]]
    assert "教师编号" in summary_labels
    assert "必填项完成率" in summary_labels


# --------------------------------------------------------------------------
# 12. 人才引进：硬条件、维度评分、证据、人工评审与候选人导入
# --------------------------------------------------------------------------
def test_recruitment_evaluation_review_and_import(client, manager_headers):
    position = client.post(
        "/api/recruitment/positions",
        json={
            "name": "自动化测试岗位",
            "college": "测试学院",
            "discipline": "人工智能",
            "research_direction": "人工智能 机器学习 大模型",
            "required_degree": "博士",
            "min_outputs": 4,
            "research_topics": "大模型 智能系统",
            "requirements_desc": "要求具备大模型研究经验",
            "weights": {"direction": 0.4, "outputs": 0.3, "projects": 0.2, "teaching": 0.1},
            "status": "open",
        },
        headers=manager_headers,
    )
    assert position.status_code == 200, position.text
    position_id = position.json()["id"]
    assert abs(sum(position.json()["weights"].values()) - 1.0) < 0.001

    candidates = {
        "甲": client.post(
            "/api/recruitment/candidates",
            json={
                "name": "测试候选人甲",
                "position_id": position_id,
                "degree": "博士",
                "graduate_school": "测试大学",
                "major": "人工智能",
                "research_direction": "人工智能 机器学习 大模型 智能系统",
                "keywords": "大模型,智能系统",
                "outputs_count": 8,
                "projects_count": 4,
                "teaching_years": 5,
                "output_evidence": "8篇代表性论文（测试证据）",
                "project_evidence": "4项项目证明",
                "teaching_evidence": "5年教学证明",
                "source": "自动化测试",
            },
            headers=manager_headers,
        ),
        "乙": client.post(
            "/api/recruitment/candidates",
            json={
                "name": "测试候选人乙",
                "position_id": position_id,
                "degree": "硕士",
                "major": "人工智能",
                "research_direction": "人工智能 机器学习 大模型",
                "outputs_count": 30,
                "projects_count": 8,
                "teaching_years": 10,
                "source": "自动化测试",
            },
            headers=manager_headers,
        ),
        "丙": client.post(
            "/api/recruitment/candidates",
            json={
                "name": "测试候选人丙",
                "position_id": position_id,
                "degree": "博士",
                "major": "人工智能",
                "research_direction": "人工智能",
                "outputs_count": 1,
                "projects_count": 0,
                "teaching_years": 0,
                "source": "自动化测试",
            },
            headers=manager_headers,
        ),
    }
    for response in candidates.values():
        assert response.status_code == 200, response.text
    candidate_ids = {key: value.json()["id"] for key, value in candidates.items()}

    evaluation = client.post(
        f"/api/recruitment/positions/{position_id}/evaluate", headers=manager_headers
    )
    assert evaluation.status_code == 200, evaluation.text
    evaluated = evaluation.json()
    assert evaluated["summary"]["total"] == 3
    assert evaluated["summary"]["hard_pass"] == 1
    assert evaluated["summary"]["hard_fail"] == 2
    assert "不自动作出录用或淘汰决定" in evaluated["summary"]["note"]

    by_name = {item["candidate"]["name"]: item for item in evaluated["items"]}
    first = by_name["测试候选人甲"]
    second = by_name["测试候选人乙"]
    third = by_name["测试候选人丙"]
    assert first["hard_pass"] is True and first["rank"] == 1
    assert second["hard_pass"] is False and second["rank"] is None
    assert third["hard_pass"] is False and third["rank"] is None
    assert "学历不满足" in second["gaps"]
    assert "成果数量不满足" in third["gaps"]
    assert "8篇代表性论文（测试证据）" in first["evidence"]
    assert "成果证据" in first["evidence"]
    assert "权重" in first["reasons"]
    assert first["review_status"] == "pending"

    reviewed = client.put(
        f"/api/recruitment/evaluations/{first['id']}/review",
        json={"review_status": "shortlisted", "review_note": "人工复核通过"},
        headers=manager_headers,
    )
    assert reviewed.status_code == 200
    assert reviewed.json()["review_status"] == "shortlisted"
    assert reviewed.json()["reviewer"] == "manager"
    assert reviewed.json()["review_note"] == "人工复核通过"

    invalid_review = client.put(
        f"/api/recruitment/evaluations/{first['id']}/review",
        json={"review_status": "auto_hired"},
        headers=manager_headers,
    )
    assert invalid_review.status_code == 422

    evaluations = client.get(
        f"/api/recruitment/positions/{position_id}/evaluations",
        headers=manager_headers,
    ).json()
    second_again = _find(
        [item for item in evaluations["items"]], candidate_id=candidate_ids["乙"]
    )
    assert second_again["review_status"] == "pending"  # 系统没有自动淘汰

    blocked = client.delete(
        f"/api/recruitment/positions/{position_id}", headers=manager_headers
    )
    assert blocked.status_code == 409
    assert "候选人" in blocked.json()["detail"]

    candidate_headers = [
        "name", "degree", "graduate_school", "major", "research_direction",
        "keywords", "outputs_count", "projects_count", "teaching_years",
        "output_evidence", "project_evidence", "teaching_evidence", "source",
        "unverified_fields",
    ]
    candidate_rows = [
        {
            "name": "测试候选人甲", "degree": "博士", "major": "人工智能",
            "research_direction": "人工智能", "source": "CSV导入",
        },
        {
            "name": "测试候选人丁", "degree": "博士", "graduate_school": "测试大学",
            "major": "人工智能", "research_direction": "人工智能 大模型",
            "keywords": "大模型", "outputs_count": "6", "projects_count": "2",
            "teaching_years": "3", "source": "CSV导入", "unverified_fields": "项目证明",
        },
    ]
    imported = client.post(
        "/api/imports/candidates",
        params={"position_id": position_id},
        files={
            "file": (
                "candidates_test.csv",
                _csv_bytes(candidate_headers, candidate_rows),
                "text/csv",
            )
        },
        headers=manager_headers,
    )
    assert imported.status_code == 200, imported.text
    result = imported.json()
    assert result["success_rows"] == 1
    assert result["duplicate_rows"] == 1

    candidate_list = client.get(
        "/api/recruitment/candidates",
        params={"position_id": position_id, "page_size": 100},
        headers=manager_headers,
    ).json()
    assert candidate_list["total"] == 4
    new_candidate = _find(candidate_list["items"], name="测试候选人丁")
    assert new_candidate["unverified_fields"] == "项目证明"


# --------------------------------------------------------------------------
# 13. 角色权限、敏感字段脱敏、设置与审计
# --------------------------------------------------------------------------
def test_permissions_masking_settings_and_audit(
    client, admin_headers, manager_headers, viewer_headers
):
    teacher = _create_teacher(
        client, manager_headers, "TEST-PERM-001",
        phone="13800138000", email="perm@example.com",
    )

    admin_view = client.get(
        "/api/teachers", params={"keyword": "TEST-PERM-001"}, headers=admin_headers
    ).json()["items"][0]
    viewer_view = client.get(
        "/api/teachers", params={"keyword": "TEST-PERM-001"}, headers=viewer_headers
    ).json()["items"][0]
    assert admin_view["phone"] == "13800138000"
    assert viewer_view["phone"] != admin_view["phone"]
    assert "***" in viewer_view["phone"]
    assert viewer_view["phone_masked"] is True
    assert "***" in viewer_view["email"]

    forbidden = client.post(
        "/api/teachers", json=_teacher_payload("TEST-PERM-002"), headers=viewer_headers
    )
    assert forbidden.status_code == 403
    assert "无权执行此操作" in forbidden.json()["detail"]

    assert client.get("/api/users", headers=viewer_headers).status_code == 403
    assert client.get("/api/users", headers=manager_headers).status_code == 403
    users = client.get("/api/users", headers=admin_headers)
    assert users.status_code == 200
    assert users.json()["total"] >= 3
    assert all("password_hash" not in user for user in users.json()["items"])

    manager_settings = client.get("/api/settings", headers=manager_headers)
    assert manager_settings.status_code == 200
    assert manager_settings.json()["can_edit"] is False
    admin_settings = client.get("/api/settings", headers=admin_headers)
    assert admin_settings.json()["can_edit"] is True

    assert client.put(
        "/api/settings/weights",
        json={"advisor_weights": {"direction": 1}},
        headers=manager_headers,
    ).status_code == 403
    weights = client.put(
        "/api/settings/weights",
        json={"advisor_weights": {
            "direction": 40, "topic": 25, "needs": 20, "capacity": 15
        }},
        headers=admin_headers,
    )
    assert weights.status_code == 200
    assert abs(sum(weights.json()["advisor_weights"].values()) - 1.0) < 0.001
    bad_weight = client.put(
        "/api/settings/weights",
        json={"advisor_weights": {"unknown": 1}},
        headers=admin_headers,
    )
    assert bad_weight.status_code == 400
    assert "未知权重维度" in bad_weight.json()["detail"]
    client.put(
        "/api/settings/weights",
        json={"advisor_weights": {
            "direction": 0.4, "topic": 0.25, "needs": 0.2, "capacity": 0.15
        }},
        headers=admin_headers,
    )

    assert client.get("/api/audit-logs", headers=viewer_headers).status_code == 403
    audit = client.get(
        "/api/audit-logs",
        params={
            "action": "create",
            "resource_type": "teacher",
            "page_size": 100,
        },
        headers=manager_headers,
    )
    assert audit.status_code == 200
    rows = audit.json()["items"]
    assert any(
        row["username"] == "manager" and row["resource_type"] == "teacher"
        for row in rows
    )
    assert all(row["detail"] is not None for row in rows)

    assert client.delete(
        f"/api/teachers/{teacher['id']}", headers=manager_headers
    ).status_code == 200


# --------------------------------------------------------------------------
# 14. 未配置 LLM Key 时本地回退、系统信息与数据标记
# --------------------------------------------------------------------------
def test_llm_fallback_and_system_info(client, viewer_headers, admin_headers):
    settings = client.get("/api/settings", headers=admin_headers).json()
    assert settings["llm"]["configured"] is False
    assert settings["llm"]["mode"] == "local"
    assert "未配置 API Key" in settings["llm"]["note"]

    first_output = client.get(
        "/api/research-outputs", params={"page_size": 1}, headers=viewer_headers
    ).json()["items"][0]
    summary = client.post(
        f"/api/teachers/{first_output['teacher_id']}/ai-summary", headers=viewer_headers
    )
    assert summary.status_code == 200, summary.text
    data = summary.json()
    assert data["mode"] == "local"
    assert "未配置 LLM_API_KEY" in data["warning"]
    assert "画像摘要" in data["text"]
    assert data["basis"]["sensitive_data_sent"] is False

    info = client.get("/api/system/info", headers=admin_headers)
    assert info.status_code == 200
    body = info.json()
    assert body["project_no"] == "B3"
    assert body["database"]["backend"] == "sqlite"
    assert body["counts"]["teachers"] > 0
    assert body["llm"]["mode"] == "local"
    assert body["demo_data"]["enabled"] is True
    assert body["demo_data"]["marker"] == "模拟数据（B3演示）"


# --------------------------------------------------------------------------
# 15. 重启后数据仍在（新应用生命周期 + 真实 SQLite 文件）
# --------------------------------------------------------------------------
def test_data_persists_across_restart(client, manager_headers):
    teacher = _create_teacher(client, manager_headers, "TEST-PERSIST-001")
    before = client.get(
        "/api/teachers", params={"page_size": 1}, headers=manager_headers
    ).json()["total"]

    with TestClient(app) as restarted:
        headers = {
            "Authorization": "Bearer "
            + restarted.post(
                "/api/auth/login",
                json={"username": "manager", "password": "manager123"},
            ).json()["access_token"]
        }
        after = restarted.get(
            "/api/teachers", params={"page_size": 1}, headers=headers
        ).json()["total"]
        assert after == before
        persisted = restarted.get(f"/api/teachers/{teacher['id']}", headers=headers)
        assert persisted.status_code == 200
        assert persisted.json()["teacher_no"] == "TEST-PERSIST-001"

    assert client.delete(
        f"/api/teachers/{teacher['id']}", headers=manager_headers
    ).status_code == 200


# --------------------------------------------------------------------------
# 16. 上传与查询错误的可理解提示
# --------------------------------------------------------------------------
def test_clear_errors_for_bad_requests(client, manager_headers):
    assert client.get("/api/teachers/999999", headers=manager_headers).status_code == 404
    assert "教师不存在" in client.get(
        "/api/teachers/999999", headers=manager_headers
    ).json()["detail"]

    invalid_page = client.get(
        "/api/teachers", params={"page": 0}, headers=manager_headers
    )
    assert invalid_page.status_code == 422
    assert invalid_page.json()["detail"].startswith("参数校验失败")

    bad_suffix = client.post(
        "/api/imports/teachers",
        files={"file": ("teachers.txt", b"abc", "text/plain")},
        headers=manager_headers,
    )
    assert bad_suffix.status_code == 400
    assert "仅支持 .csv 或 .xlsx 文件" in bad_suffix.json()["detail"]

    empty = client.post(
        "/api/imports/teachers",
        files={"file": ("empty.csv", b"", "text/csv")},
        headers=manager_headers,
    )
    assert empty.status_code == 400
    assert "上传文件为空" in empty.json()["detail"]
