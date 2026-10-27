"""Database-backed statistics for the management dashboard and talent review."""
import json
import re
from collections import Counter, defaultdict
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from ..models import (
    ProjectMember,
    ResearchOutput,
    ResearchProject,
    ServiceRecord,
    SystemSetting,
    Teacher,
    TeachingRecord,
)
from ..utils import split_keywords
from .directory import group_disciplines_by_college

TITLE_ORDER = ["教授", "副教授", "讲师", "助教"]
DEGREE_ORDER = ["博士", "硕士", "学士", "其他"]
OUTPUT_TYPE_LABELS = {"paper": "论文", "patent": "专利", "book": "著作", "report": "研究报告"}
SERVICE_TYPE_LABELS = {"award": "获奖", "academic_post": "学术任职", "social_service": "社会服务"}
AGE_BANDS = ["30岁及以下", "31-40岁", "41-50岁", "51-60岁", "60岁以上", "未填写"]


def _year(value):
    if value is None:
        return None
    if hasattr(value, "year"):
        return value.year
    try:
        return int(str(value)[:4])
    except (ValueError, TypeError):
        return None


def _in_years(value, year_from, year_to):
    year = _year(value)
    if year is None:
        return year_from is None and year_to is None
    if year_from and year < year_from:
        return False
    if year_to and year > year_to:
        return False
    return True


def get_teachers(db: Session, college=None, discipline=None, title=None, status=None) -> list[Teacher]:
    query = db.query(Teacher)
    if college:
        query = query.filter(Teacher.college == college)
    if discipline:
        query = query.filter(Teacher.discipline == discipline)
    if title:
        query = query.filter(Teacher.title == title)
    if status:
        query = query.filter(Teacher.status == status)
    return query.order_by(Teacher.id).all()


def get_options(db: Session) -> dict:
    teachers = db.query(Teacher).all()
    years = set()
    for row in db.query(ResearchOutput.publish_date).all():
        y = _year(row[0])
        if y:
            years.add(y)
    for row in db.query(ResearchProject.start_date).all():
        y = _year(row[0])
        if y:
            years.add(y)
    return {
        "colleges": sorted({t.college for t in teachers if t.college}),
        "disciplines": sorted({t.discipline for t in teachers if t.discipline}),
        "disciplines_by_college": group_disciplines_by_college(teachers),
        "titles": [t for t in TITLE_ORDER if any(x.title == t for x in teachers)],
        "statuses": sorted({t.status for t in teachers if t.status}),
        "years": sorted(years),
    }


def data_quality(db: Session, teachers: list[Teacher]) -> dict:
    teacher_ids = [t.id for t in teachers]
    missing_fields = Counter()
    check_fields = [
        ("degree", "学历"), ("title", "职称"), ("research_direction", "研究方向"),
        ("research_keywords", "研究关键词"), ("phone", "联系方式"), ("email", "邮箱"),
        ("hire_date", "入职时间"),
    ]
    for teacher in teachers:
        for field, _ in check_fields:
            if not getattr(teacher, field, None):
                missing_fields[field] += 1

    name_groups = defaultdict(list)
    for teacher in teachers:
        name_groups[(teacher.name, teacher.college)].append(teacher.teacher_no)
    duplicate_teachers = [
        {"name": name, "college": college, "teacher_nos": nos, "count": len(nos)}
        for (name, college), nos in name_groups.items() if len(nos) > 1
    ]

    duplicate_outputs = []
    if teacher_ids:
        output_groups = defaultdict(list)
        teacher_name = {t.id: t.name for t in teachers}
        for row in db.query(ResearchOutput).filter(ResearchOutput.teacher_id.in_(teacher_ids)).all():
            key = (row.teacher_id, re.sub(r"\s+", "", (row.title or "").lower()))
            output_groups[key].append(row)
        for (teacher_id, _), rows in output_groups.items():
            if len(rows) > 1:
                duplicate_outputs.append({
                    "teacher_id": teacher_id,
                    "teacher_name": teacher_name.get(teacher_id),
                    "title": rows[0].title,
                    "count": len(rows),
                    "ids": [r.id for r in rows],
                })

    pending_teachers = [
        {
            "id": t.id, "teacher_no": t.teacher_no, "name": t.name,
            "college": t.college, "verification_status": t.verification_status,
            "data_source": t.data_source,
        }
        for t in teachers if t.verification_status != "verified"
    ]

    no_output_ids = set()
    if teacher_ids:
        no_output_ids = set(teacher_ids) - {
            r[0] for r in db.query(ResearchOutput.teacher_id)
            .filter(ResearchOutput.teacher_id.in_(teacher_ids)).distinct().all()
        }

    return {
        "missing_field_counts": [
            {"field": field, "label": label, "count": missing_fields.get(field, 0)}
            for field, label in check_fields
        ],
        "duplicate_teacher_count": len(duplicate_teachers),
        "duplicate_teachers": duplicate_teachers[:20],
        "duplicate_output_count": len(duplicate_outputs),
        "duplicate_outputs": duplicate_outputs[:20],
        "pending_verification_count": len(pending_teachers),
        "pending_verification": pending_teachers[:20],
        "teachers_without_outputs": len(no_output_ids),
        "teachers_without_outputs_sample": [
            {"id": t.id, "teacher_no": t.teacher_no, "name": t.name, "college": t.college}
            for t in teachers if t.id in no_output_ids
        ][:20],
        "note": "数据质量检查基于档案与成果表实际记录，仅反映数据完整性/一致性，不代表成果真实性。",
    }


def dashboard(db: Session, college=None, discipline=None, title=None,
              year_from=None, year_to=None, status=None) -> dict:
    teachers = get_teachers(db, college, discipline, title, status)
    teacher_ids = [t.id for t in teachers]
    id_set = set(teacher_ids)

    if teacher_ids:
        outputs = db.query(ResearchOutput).filter(
            ResearchOutput.teacher_id.in_(teacher_ids)
        ).order_by(ResearchOutput.publish_date).all()
        output_rows = [o for o in outputs if _in_years(o.publish_date, year_from, year_to)]
        teaching = db.query(TeachingRecord).filter(
            TeachingRecord.teacher_id.in_(teacher_ids)
        ).all()
        teaching = [t for t in teaching if _in_years(t.academic_year, year_from, year_to)]
        services = db.query(ServiceRecord).filter(
            ServiceRecord.teacher_id.in_(teacher_ids)
        ).all()
        services = [s for s in services if _in_years(s.start_date, year_from, year_to)]

        member_project_ids = {
            r[0] for r in db.query(ProjectMember.project_id)
            .filter(ProjectMember.teacher_id.in_(teacher_ids)).all()
        }
        project_query = db.query(ResearchProject)
        if member_project_ids:
            project_query = project_query.filter(
                (ResearchProject.principal_teacher_id.in_(teacher_ids))
                | (ResearchProject.id.in_(member_project_ids))
            )
        else:
            project_query = project_query.filter(ResearchProject.principal_teacher_id.in_(teacher_ids))
        projects = [p for p in project_query.all() if _in_years(p.start_date, year_from, year_to)]
        principal_projects = [p for p in projects if p.principal_teacher_id in id_set]
    else:
        output_rows, teaching, services, projects, principal_projects = [], [], [], [], []

    output_year = Counter()
    output_type = Counter()
    output_level = Counter()
    for item in output_rows:
        y = _year(item.publish_date)
        if y:
            output_year[y] += 1
        output_type[item.output_type] += 1
        if item.level:
            output_level[item.level] += 1

    project_year = Counter()
    project_funding = 0.0
    for p in projects:
        y = _year(p.start_date)
        if y:
            project_year[y] += 1
        project_funding += p.funding or 0

    teaching_year = Counter()
    teaching_hours = 0.0
    eval_scores = []
    teaching_achievements = 0
    for t in teaching:
        y = _year(t.academic_year)
        if y:
            teaching_year[y] += t.hours or 0
        teaching_hours += t.hours or 0
        if t.evaluation_score is not None:
            eval_scores.append(t.evaluation_score)
        if t.achievement:
            teaching_achievements += 1

    service_type = Counter(s.service_type for s in services)

    age_dist = Counter()
    current_year = date.today().year
    for t in teachers:
        if not t.birth_year:
            age_dist["未填写"] += 1
            continue
        age = current_year - t.birth_year
        if age <= 30:
            age_dist["30岁及以下"] += 1
        elif age <= 40:
            age_dist["31-40岁"] += 1
        elif age <= 50:
            age_dist["41-50岁"] += 1
        elif age <= 60:
            age_dist["51-60岁"] += 1
        else:
            age_dist["60岁以上"] += 1

    keyword_counter = Counter()
    for t in teachers:
        for kw in split_keywords(t.research_keywords):
            keyword_counter[kw] += 1

    return {
        "filters": {
            "college": college, "discipline": discipline, "title": title,
            "year_from": year_from, "year_to": year_to, "status": status,
        },
        "total_teachers": len(teachers),
        "active_teachers": sum(1 for t in teachers if t.status == "active"),
        "advisor_qualified_count": sum(1 for t in teachers if t.advisor_qualified),
        "open_enrollment_count": sum(
            1 for t in teachers if t.advisor_qualified and t.enrollment_status == "open"
        ),
        "college_distribution": [
            {"name": k, "value": v}
            for k, v in Counter(t.college for t in teachers).most_common()
        ],
        "discipline_distribution": [
            {"name": k, "value": v}
            for k, v in Counter(t.discipline for t in teachers).most_common()
        ],
        "title_distribution": [
            {"name": k, "value": v}
            for k, v in Counter(t.title for t in teachers).most_common()
        ],
        "degree_distribution": [
            {"name": k, "value": v}
            for k, v in Counter(t.degree or "未填写" for t in teachers).most_common()
        ],
        "age_distribution": [
            {"name": band, "value": age_dist.get(band, 0)} for band in AGE_BANDS
        ],
        "status_distribution": [
            {"name": k, "value": v}
            for k, v in Counter(t.status for t in teachers).most_common()
        ],
        "output_total": len(output_rows),
        "output_by_year": [
            {"year": str(y), "value": output_year[y]} for y in sorted(output_year)
        ],
        "output_by_type": [
            {"name": OUTPUT_TYPE_LABELS.get(k, k), "value": v}
            for k, v in output_type.most_common()
        ],
        "output_by_level": [
            {"name": k, "value": v} for k, v in output_level.most_common()
        ],
        "project_total": len(projects),
        "principal_project_total": len(principal_projects),
        "project_funding_total": round(project_funding, 2),
        "project_by_year": [
            {"year": str(y), "value": project_year[y]} for y in sorted(project_year)
        ],
        "teaching_hours_total": round(teaching_hours, 1),
        "teaching_achievement_count": teaching_achievements,
        "avg_evaluation_score": round(sum(eval_scores) / len(eval_scores), 2) if eval_scores else None,
        "teaching_by_year": [
            {"year": str(y), "value": round(teaching_year[y], 1)}
            for y in sorted(teaching_year)
        ],
        "service_total": len(services),
        "service_by_type": [
            {"name": SERVICE_TYPE_LABELS.get(k, k), "value": v}
            for k, v in service_type.most_common()
        ],
        "top_keywords": [{"name": k, "value": v} for k, v in keyword_counter.most_common(20)],
        "data_quality": data_quality(db, teachers),
        "age_note": "年龄梯队由模拟档案中的出生年份计算；真实环境中需确认数据合规与使用授权。",
        "stats_scope": "所有指标均由数据库实时查询计算，筛选条件同时作用于教师及其关联记录。",
    }


def talent_review(db: Session, college=None, discipline=None, title=None,
                  year_from=None, year_to=None) -> dict:
    base = dashboard(db, college, discipline, title, year_from, year_to)
    teachers = get_teachers(db, college, discipline, title)
    title_by_college = defaultdict(Counter)
    degree_by_college = defaultdict(Counter)
    for t in teachers:
        title_by_college[t.college][t.title] += 1
        degree_by_college[t.college][t.degree or "未填写"] += 1

    plan_raw = db.query(SystemSetting).filter(SystemSetting.key == "talent_plan").first()
    plan = []
    if plan_raw and plan_raw.value:
        try:
            plan = json.loads(plan_raw.value)
        except json.JSONDecodeError:
            plan = []

    current_count = Counter((t.college, t.discipline) for t in teachers)
    gaps = []
    for item in plan:
        college_name = item.get("college")
        discipline_name = item.get("discipline")
        target = int(item.get("target_count") or 0)
        current = current_count.get((college_name, discipline_name), 0)
        gaps.append({
            "college": college_name,
            "discipline": discipline_name,
            "target_count": target,
            "current_count": current,
            "gap": target - current,
            "note": item.get("note", ""),
        })

    return {
        "overview": {
            key: base[key] for key in [
                "total_teachers", "active_teachers", "output_total", "project_total",
                "teaching_hours_total", "advisor_qualified_count", "college_distribution",
                "discipline_distribution", "title_distribution", "degree_distribution",
                "output_by_year", "project_by_year", "top_keywords", "avg_evaluation_score",
                "stats_scope", "age_note",
            ]
        },
        "title_structure_by_college": [
            {
                "college": college_name,
                "total": sum(counts.values()),
                **{title_name: counts.get(title_name, 0) for title_name in TITLE_ORDER},
            }
            for college_name, counts in sorted(title_by_college.items())
        ],
        "degree_structure_by_college": [
            {
                "college": college_name,
                "total": sum(counts.values()),
                **{degree_name: counts.get(degree_name, 0) for degree_name in DEGREE_ORDER},
            }
            for college_name, counts in sorted(degree_by_college.items())
        ],
        "talent_plan": plan,
        "plan_gaps": gaps,
        "plan_note": (
            "人才缺口依据系统设置中管理员配置的人才规划目标计算，非模型推断；"
            "未配置规划时缺口为空。"
        ),
        "data_quality": base["data_quality"],
        "filters": base["filters"],
        "generated_at": base["generated_at"] if "generated_at" in base else None,
    }


def build_report_workbook(review: dict, label: str = "") -> BytesIO:
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5597")

    def style_header(ws):
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill

    overview = review["overview"]
    ws = wb.active
    ws.title = "人才盘点概览"
    ws.append(["指标", "数值"])
    rows = [
        ("筛选条件", label or "全部"),
        ("教师总人数", overview["total_teachers"]),
        ("在职教师数", overview["active_teachers"]),
        ("科研成果总数", overview["output_total"]),
        ("参与项目总数", overview["project_total"]),
        ("教学工作量合计（学时）", overview["teaching_hours_total"]),
        ("具有导师资格人数", overview["advisor_qualified_count"]),
        ("平均教学评价", overview["avg_evaluation_score"] if overview["avg_evaluation_score"] is not None else "无数据"),
        ("统计说明", overview["stats_scope"]),
    ]
    for row in rows:
        ws.append(list(row))
    style_header(ws)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 80

    ws2 = wb.create_sheet("学院职称结构")
    ws2.append(["学院", "合计"] + TITLE_ORDER)
    for row in review["title_structure_by_college"]:
        ws2.append([row["college"], row["total"]] + [row.get(t, 0) for t in TITLE_ORDER])
    style_header(ws2)
    for col in "ABCDEF":
        ws2.column_dimensions[col].width = 14

    ws3 = wb.create_sheet("学院学历结构")
    ws3.append(["学院", "合计"] + DEGREE_ORDER)
    for row in review["degree_structure_by_college"]:
        ws3.append([row["college"], row["total"]] + [row.get(d, 0) for d in DEGREE_ORDER])
    style_header(ws3)
    for col in "ABCDEF":
        ws3.column_dimensions[col].width = 14

    ws4 = wb.create_sheet("年度成果")
    ws4.append(["年份", "成果数"])
    for row in overview["output_by_year"]:
        ws4.append([row["year"], row["value"]])
    style_header(ws4)
    ws4.column_dimensions["A"].width = 12
    ws4.column_dimensions["B"].width = 12

    ws5 = wb.create_sheet("人才规划缺口")
    ws5.append(["学院", "学科", "规划目标", "现有教师", "缺口", "备注"])
    for row in review["plan_gaps"]:
        ws5.append([
            row["college"], row["discipline"], row["target_count"],
            row["current_count"], row["gap"], row["note"],
        ])
    if not review["plan_gaps"]:
        ws5.append(["未配置人才规划", "", "", "", "", review["plan_note"]])
    style_header(ws5)
    for col, width in zip("ABCDEF", [18, 24, 12, 12, 10, 40]):
        ws5.column_dimensions[col].width = width

    quality = review["data_quality"]
    ws6 = wb.create_sheet("数据质量")
    ws6.append(["检查项", "数量"])
    ws6.append(["档案缺字段合计", sum(item["count"] for item in quality["missing_field_counts"])])
    ws6.append(["疑似重复教师组", quality["duplicate_teacher_count"]])
    ws6.append(["疑似重复成果组", quality["duplicate_output_count"]])
    ws6.append(["待核验教师", quality["pending_verification_count"]])
    ws6.append(["无成果记录教师", quality["teachers_without_outputs"]])
    ws6.append(["说明", quality["note"]])
    style_header(ws6)
    ws6.column_dimensions["A"].width = 26
    ws6.column_dimensions["B"].width = 80

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
