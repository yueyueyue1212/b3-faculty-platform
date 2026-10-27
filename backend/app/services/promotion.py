"""职称申报材料自动归集与一致性/完整性检查（不验证成果真伪）。"""
import re
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from ..models import (
    PromotionBatch,
    PromotionMaterial,
    PromotionRequirement,
    ProjectMember,
    ResearchOutput,
    ResearchProject,
    ServiceRecord,
    Teacher,
    TeachingRecord,
)
from ..utils import split_keywords

CATEGORY_LABELS = {
    "paper": "论文/著作/专利",
    "project": "科研项目",
    "award": "获奖",
    "teaching": "教学",
    "service": "学术任职/社会服务",
}

STATUS_LABELS = {
    "matched": "已匹配",
    "missing": "缺项",
    "pending_verification": "待核验",
    "duplicate": "重复",
    "conflict": "信息冲突",
    "confirmed": "人工确认",
    "rejected": "人工驳回",
}

MANUAL_STATUSES = ("confirmed", "rejected")


def _norm_title(text: str) -> str:
    return re.sub(r"[\s\-—_，,。.;；:：()（）]", "", (text or "").lower())


def _in_window(day: date | None, window_start: int, window_end_year: int) -> bool:
    if day is None:
        return True  # missing date -> included but flagged for verification
    return window_start <= day.year <= window_end_year


def requirement_matches(db: Session, teacher: Teacher, req: PromotionRequirement,
                        batch: PromotionBatch) -> list[dict]:
    start_year = (batch.year or date.today().year) - (req.window_years or 5)
    end_year = batch.year or date.today().year
    keywords = [k.lower() for k in split_keywords(req.keywords)]
    matches: list[dict] = []

    def keyword_ok(text: str) -> bool:
        if not keywords:
            return True
        haystack = (text or "").lower()
        return any(k in haystack for k in keywords)

    if req.category in ("paper", "project", "teaching", "award", "service"):
        pass

    if req.category == "paper":
        rows = db.query(ResearchOutput).filter(ResearchOutput.teacher_id == teacher.id).all()
        for row in rows:
            if not _in_window(row.publish_date, start_year, end_year):
                continue
            if not keyword_ok(row.title):
                continue
            matches.append({
                "title": row.title,
                "output_type": row.output_type,
                "ref": f"output:{row.id}",
                "date": row.publish_date,
                "source": row.data_source,
            })

    elif req.category == "project":
        member_ids = {
            r[0] for r in db.query(ProjectMember.project_id)
            .filter(ProjectMember.teacher_id == teacher.id).all()
        }
        query = db.query(ResearchProject)
        if req.principal_only:
            query = query.filter(ResearchProject.principal_teacher_id == teacher.id)
        else:
            if member_ids:
                query = query.filter(
                    (ResearchProject.principal_teacher_id == teacher.id)
                    | (ResearchProject.id.in_(member_ids))
                )
            else:
                query = query.filter(ResearchProject.principal_teacher_id == teacher.id)
        for row in query.all():
            check_date = row.start_date or row.end_date
            if not _in_window(check_date, start_year, end_year):
                continue
            if not keyword_ok(row.name):
                continue
            matches.append({
                "title": row.name,
                "output_type": "project",
                "ref": f"project:{row.id}",
                "date": check_date,
                "source": row.data_source,
            })

    elif req.category == "award":
        for row in db.query(ServiceRecord).filter(
            ServiceRecord.teacher_id == teacher.id, ServiceRecord.service_type == "award"
        ).all():
            if not _in_window(row.start_date, start_year, end_year) or not keyword_ok(row.name):
                continue
            matches.append({
                "title": row.name,
                "output_type": "award",
                "ref": f"service:{row.id}",
                "date": row.start_date,
                "source": row.data_source,
            })
        for row in db.query(TeachingRecord).filter(
            TeachingRecord.teacher_id == teacher.id, TeachingRecord.achievement.isnot(None)
        ).all():
            if not row.achievement or not keyword_ok(row.achievement):
                continue
            year = None
            try:
                year = int(str(row.academic_year)[:4])
            except (TypeError, ValueError):
                pass
            if year is not None and not (start_year <= year <= end_year):
                continue
            matches.append({
                "title": f"教学成果：{row.achievement}",
                "output_type": "teaching_award",
                "ref": f"teaching:{row.id}",
                "date": date(year, 12, 31) if year else None,
                "source": row.data_source,
            })

    elif req.category == "teaching":
        for row in db.query(TeachingRecord).filter(TeachingRecord.teacher_id == teacher.id).all():
            try:
                year = int(str(row.academic_year)[:4])
            except (TypeError, ValueError):
                year = None
            if year is not None and not (start_year <= year <= end_year):
                continue
            if not keyword_ok(row.course_name):
                continue
            matches.append({
                "title": row.course_name,
                "output_type": "teaching",
                "ref": f"teaching:{row.id}",
                "date": date(year, 12, 31) if year else None,
                "source": row.data_source,
            })

    elif req.category == "service":
        for row in db.query(ServiceRecord).filter(
            ServiceRecord.teacher_id == teacher.id,
            ServiceRecord.service_type.in_(["academic_post", "social_service"]),
        ).all():
            if not _in_window(row.start_date, start_year, end_year) or not keyword_ok(row.name):
                continue
            matches.append({
                "title": row.name,
                "output_type": row.service_type,
                "ref": f"service:{row.id}",
                "date": row.start_date,
                "source": row.data_source,
            })

    return matches


def gather_materials(db: Session, batch: PromotionBatch, teacher: Teacher) -> dict:
    requirements = db.query(PromotionRequirement).filter(
        PromotionRequirement.batch_id == batch.id
    ).order_by(PromotionRequirement.id).all()

    manual = db.query(PromotionMaterial).filter(
        PromotionMaterial.batch_id == batch.id,
        PromotionMaterial.teacher_id == teacher.id,
        PromotionMaterial.status.in_(MANUAL_STATUSES),
    ).all()
    manual_titles = {(m.requirement_id, _norm_title(m.title)) for m in manual}

    db.query(PromotionMaterial).filter(
        PromotionMaterial.batch_id == batch.id,
        PromotionMaterial.teacher_id == teacher.id,
        ~PromotionMaterial.status.in_(MANUAL_STATUSES),
    ).delete(synchronize_session=False)
    db.flush()

    created = 0
    for req in requirements:
        matches = requirement_matches(db, teacher, req, batch)
        seen = {}
        valid_count = 0
        for item in matches:
            key = _norm_title(item["title"])
            if (req.id, key) in manual_titles:
                continue
            status = "matched"
            if key in seen:
                status = "duplicate"
            elif not item["source"] or "待核验" in (item["source"] or ""):
                status = "pending_verification"
            elif item["date"] is None:
                status = "pending_verification"
            elif batch.deadline and item["date"] > batch.deadline:
                status = "conflict"
            if status not in ("duplicate", "conflict"):
                valid_count += 1
            seen[key] = True
            db.add(PromotionMaterial(
                batch_id=batch.id,
                teacher_id=teacher.id,
                requirement_id=req.id,
                category=req.category,
                title=item["title"],
                output_type=item["output_type"],
                output_ref=item["ref"],
                matched_count=1,
                status=status,
            ))
            created += 1

        manual_count = sum(
            1 for m in manual if m.requirement_id == req.id and m.status == "confirmed"
        )
        total_valid = valid_count + manual_count
        if total_valid < (req.min_count or 1):
            already_manual = any(
                m.requirement_id == req.id and m.title.startswith("缺项") for m in manual
            )
            if not already_manual:
                db.add(PromotionMaterial(
                    batch_id=batch.id,
                    teacher_id=teacher.id,
                    requirement_id=req.id,
                    category=req.category,
                    title=f"缺项：{req.name}（已匹配 {total_valid}/{req.min_count}）",
                    output_type="missing",
                    matched_count=total_valid,
                    status="missing",
                    note="系统自动识别，请核实数据或人工补充材料",
                ))
                created += 1

    db.commit()
    return {"created": created, "teacher_id": teacher.id, "batch_id": batch.id}


def promotion_summary(db: Session, batch: PromotionBatch, teacher_id: int) -> dict:
    requirements = db.query(PromotionRequirement).filter(
        PromotionRequirement.batch_id == batch.id
    ).order_by(PromotionRequirement.id).all()
    materials = db.query(PromotionMaterial).filter(
        PromotionMaterial.batch_id == batch.id,
        PromotionMaterial.teacher_id == teacher_id,
    ).order_by(PromotionMaterial.requirement_id, PromotionMaterial.id).all()

    counts = {key: 0 for key in STATUS_LABELS}
    for m in materials:
        counts[m.status] = counts.get(m.status, 0) + 1

    rows = []
    missing_mandatory = []
    mandatory_total = 0
    mandatory_satisfied = 0
    for req in requirements:
        req_materials = [m for m in materials if m.requirement_id == req.id]
        present = sum(
            1 for m in req_materials
            if m.status in ("matched", "pending_verification", "confirmed")
        )
        pending = sum(1 for m in req_materials if m.status == "pending_verification")
        duplicates = sum(1 for m in req_materials if m.status == "duplicate")
        conflicts = sum(1 for m in req_materials if m.status == "conflict")
        satisfied = present >= (req.min_count or 1)
        if req.mandatory:
            mandatory_total += 1
            if satisfied:
                mandatory_satisfied += 1
            else:
                missing_mandatory.append({
                    "requirement_id": req.id,
                    "name": req.name,
                    "category": req.category,
                    "category_label": CATEGORY_LABELS.get(req.category, req.category),
                    "min_count": req.min_count,
                    "present": present,
                })
        rows.append({
            "requirement_id": req.id,
            "name": req.name,
            "category": req.category,
            "category_label": CATEGORY_LABELS.get(req.category, req.category),
            "requirement_desc": req.requirement_desc,
            "min_count": req.min_count,
            "mandatory": req.mandatory,
            "present": present,
            "pending": pending,
            "duplicates": duplicates,
            "conflicts": conflicts,
            "satisfied": satisfied,
        })

    completion = round(mandatory_satisfied / mandatory_total * 100, 1) if mandatory_total else 100.0
    return {
        "batch": {
            "id": batch.id,
            "name": batch.name,
            "year": batch.year,
            "title_type": batch.title_type,
            "college": batch.college,
            "status": batch.status,
            "deadline": batch.deadline.isoformat() if batch.deadline else None,
        },
        "teacher_id": teacher_id,
        "counts": counts,
        "status_labels": STATUS_LABELS,
        "requirements": rows,
        "materials": [
            {
                "id": m.id,
                "requirement_id": m.requirement_id,
                "category": m.category,
                "category_label": CATEGORY_LABELS.get(m.category, m.category),
                "title": m.title,
                "output_type": m.output_type,
                "output_ref": m.output_ref,
                "matched_count": m.matched_count,
                "status": m.status,
                "status_label": STATUS_LABELS.get(m.status, m.status),
                "note": m.note,
                "updated_at": m.updated_at.isoformat() if m.updated_at else None,
            }
            for m in materials
        ],
        "missing_mandatory": missing_mandatory,
        "completion": completion,
        "disclaimer": "材料检查仅比对已归集数据的一致性、完整性与重复性，不验证成果真伪。",
    }


def build_promotion_workbook(summary: dict, teacher: Teacher) -> BytesIO:
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5597")

    ws = wb.active
    ws.title = "申报汇总"
    ws.append(["项目", "内容"])
    rows = [
        ("申报批次", summary["batch"]["name"]),
        ("职称类型", summary["batch"]["title_type"]),
        ("学院", summary["batch"]["college"] or "不限"),
        ("教师编号", teacher.teacher_no),
        ("教师姓名", teacher.name),
        ("必填项完成率", f"{summary['completion']}%"),
        ("缺项数量", len([m for m in summary["materials"] if m["status"] == "missing"])),
        ("待核验数量", summary["counts"].get("pending_verification", 0)),
        ("重复数量", summary["counts"].get("duplicate", 0)),
        ("冲突数量", summary["counts"].get("conflict", 0)),
        ("说明", summary["disclaimer"]),
    ]
    for row in rows:
        ws.append(row)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 80

    ws2 = wb.create_sheet("要求达成情况")
    ws2.append(["类别", "申报要求", "最低数量", "已匹配", "待核验", "重复", "冲突", "是否满足", "是否必填"])
    for req in summary["requirements"]:
        ws2.append([
            req["category_label"], req["name"], req["min_count"], req["present"],
            req["pending"], req["duplicates"], req["conflicts"],
            "是" if req["satisfied"] else "否", "是" if req["mandatory"] else "否",
        ])
    for cell in ws2[1]:
        cell.font = header_font
        cell.fill = header_fill
    for col, width in zip("ABCDEFGHI", [14, 40, 10, 10, 10, 10, 10, 10, 10]):
        ws2.column_dimensions[col].width = width

    ws3 = wb.create_sheet("材料明细")
    ws3.append(["类别", "材料名称", "引用编号", "数量", "状态", "备注"])
    for m in summary["materials"]:
        ws3.append([
            m["category_label"], m["title"], m["output_ref"] or "", m["matched_count"],
            m["status_label"], m["note"] or "",
        ])
    for cell in ws3[1]:
        cell.font = header_font
        cell.fill = header_fill
    for col, width in zip("ABCDEF", [14, 50, 18, 8, 12, 40]):
        ws3.column_dimensions[col].width = width

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
